# -*- coding: utf-8 -*-
"""
tools.py · v3.0.0 本地 MCP 风格工具注册表
----------------------------------------------------------------
MCP 兼容描述符结构：{"name","description","inputSchema"(JSON Schema),"handler"}
将来接入外部标准 MCP server：用官方 mcp SDK 的 ClientSession.list_tools() 拿到
tool.inputSchema（与本表同构），包一层 call_tool handler 混入 REGISTRY 即可，tool loop 零改动。
TODO(Worker): 工具清单将来可由 Worker/数据库下发合并。
"""
import json
import os
import csv
import io
import re
import urllib.request
import datetime

# ---------- 路径 ----------

def _shell_folder(name):
    try:
        import winreg
        k = winreg.OpenKey(winreg.HKEY_CURRENT_USER,
                           r'Software\Microsoft\Windows\CurrentVersion\Explorer\Shell Folders')
        v, _ = winreg.QueryValueEx(k, name)
        winreg.CloseKey(k)
        return v
    except Exception:
        return None

def desktop_dir():
    for cand in [_shell_folder('Desktop'),
                 os.path.join(os.path.expanduser('~'), 'Desktop'),
                 os.path.join(os.path.expanduser('~'), 'OneDrive', 'Desktop')]:
        if cand and os.path.isdir(cand):
            return cand
    return os.path.expanduser('~')

def documents_dir():
    d = _shell_folder('Personal')
    return d if d and os.path.isdir(d) else os.path.join(os.path.expanduser('~'), 'Documents')

def export_dir():
    d = os.path.join(desktop_dir(), 'chatgpt-ui-exports')
    os.makedirs(d, exist_ok=True)
    return d

def _safe_name(name, default_ext):
    name = re.sub(r'[\\/:*?"<>|]', '_', (name or '').strip()) or 'untitled'
    if not name.lower().endswith(default_ext):
        name += default_ext
    return name

# ---------- Word ----------

def create_word(filename, title, content):
    """content: markdown-lite 文本。# ## 标题 / - 列表 / 1. 编号 / 普通段落。"""
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    if title:
        doc.add_heading(title, level=0)
    for line in (content or '').splitlines():
        line = line.rstrip()
        if not line.strip():
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif re.match(r'^[-*] ', line):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif re.match(r'^\d+\. ', line):
            doc.add_paragraph(re.sub(r'^\d+\. ', '', line), style='List Number')
        else:
            doc.add_paragraph(line)
    path = os.path.join(export_dir(), _safe_name(filename, '.docx'))
    doc.save(path)
    return path

# ---------- Excel ----------

def create_excel(filename, rows, sheet_name='Sheet1'):
    """rows: [[...], [...]] 二维数组，首行当表头。也接受 csv_text。"""
    import openpyxl
    from openpyxl.styles import Font
    if isinstance(rows, str):
        rows = list(csv.reader(io.StringIO(rows)))
    rows = [list(map(lambda c: c if not isinstance(c, str) else c, r)) for r in rows]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (sheet_name or 'Sheet1')[:31]
    for r in rows:
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    path = os.path.join(export_dir(), _safe_name(filename, '.xlsx'))
    wb.save(path)
    return path

# ---------- PPT ----------

def create_ppt(filename, title, slides):
    """slides: [{"title": "...", "bullets": ["...", ...]}, ...]"""
    from pptx import Presentation
    from pptx.util import Pt as PPt
    prs = Presentation()
    # 标题页
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title or filename
    s.placeholders[1].text = datetime.date.today().isoformat() + ' · 由 ChatGPT UI 桌面版生成'
    for sl in slides or []:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = str(sl.get('title', ''))[:80]
        body = s.placeholders[1]
        body.text.clear() if hasattr(body.text, 'clear') else None
        bullets = sl.get('bullets') or []
        first = True
        for b in bullets[:12]:
            p = body.text_frame.paragraphs[0] if first else body.text_frame.add_paragraph()
            first = False
            p.text = str(b)
            p.font.size = PPt(18)
    path = os.path.join(export_dir(), _safe_name(filename, '.pptx'))
    prs.save(path)
    return path

# ---------- 文件系统 ----------

def list_desktop_files(ext=None):
    d = desktop_dir()
    items = []
    for n in sorted(os.listdir(d))[:200]:
        p = os.path.join(d, n)
        if os.path.isfile(p):
            if ext and not n.lower().endswith(str(ext).lower()):
                continue
            st = os.stat(p)
            items.append({'name': n, 'size_kb': round(st.st_size / 1024, 1),
                          'mtime': datetime.datetime.fromtimestamp(st.st_mtime).strftime('%Y-%m-%d %H:%M')})
    return {'dir': d, 'files': items[:50]}

OPEN_WHITELIST = None  # 由 list_desktop_files/export_dir 动态判断

def open_file(filename):
    d = desktop_dir()
    e = export_dir()
    for base in (d, e):
        p = os.path.join(base, filename)
        if os.path.isfile(p) and os.path.abspath(p).startswith(os.path.abspath(base)):
            os.startfile(p)
            return {'opened': p}
    raise FileNotFoundError('文件不在桌面或导出目录中：' + filename)

# ---------- 网页抓取 ----------

_TAG_RE = re.compile(r'<(script|style)[\s\S]*?</\1>|<[^>]+>', re.I)

def fetch_webpage(url, max_chars=8000):
    if not re.match(r'^https?://', url or ''):
        raise ValueError('URL 必须以 http(s):// 开头')
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'})
    with urllib.request.urlopen(req, timeout=30) as r:
        raw = r.read(2 * 1024 * 1024)
    try:
        html = raw.decode(r.headers.get_content_charset() or 'utf-8', 'ignore')
    except Exception:
        html = raw.decode('utf-8', 'ignore')
    m = re.search(r'<title[^>]*>(.*?)</title>', html, re.I | re.S)
    title = re.sub(r'\s+', ' ', m.group(1)).strip() if m else url
    text = _TAG_RE.sub(' ', html)
    text = re.sub(r'\s+', ' ', text).strip()
    return {'title': title, 'url': url, 'text': text[:max_chars]}

# ---------- 注册表 ----------

def _h(fn):
    return lambda args: fn(**args)

REGISTRY = [
    {
        'name': 'create_word',
        'description': '生成 Word 文档(.docx)并保存到桌面 chatgpt-ui-exports 目录。content 支持 # 标题、- 列表、编号列表、普通段落。',
        'inputSchema': {'type': 'object', 'properties': {
            'filename': {'type': 'string', 'description': '文件名，不含扩展名'},
            'title': {'type': 'string', 'description': '文档大标题'},
            'content': {'type': 'string', 'description': '文档正文（markdown-lite 文本）'}},
            'required': ['filename', 'title', 'content']},
        'handler': _h(lambda filename, title, content: create_word(filename, title, content)),
    },
    {
        'name': 'create_excel',
        'description': '生成 Excel 表格(.xlsx)并保存到桌面 chatgpt-ui-exports 目录。rows 为二维数组，首行是表头。',
        'inputSchema': {'type': 'object', 'properties': {
            'filename': {'type': 'string'},
            'sheet_name': {'type': 'string'},
            'rows': {'type': 'array', 'items': {'type': 'array'}, 'description': '二维数组，如 [["姓名","成绩"],["张三",95]]'}},
            'required': ['filename', 'rows']},
        'handler': _h(lambda filename, rows, sheet_name='Sheet1': create_excel(filename, rows, sheet_name)),
    },
    {
        'name': 'create_ppt',
        'description': '生成 PPT 演示文稿(.pptx)并保存到桌面 chatgpt-ui-exports 目录。slides 为每页 {title, bullets[]}。',
        'inputSchema': {
            'type': 'object',
            'properties': {
                'filename': {'type': 'string'},
                'title': {'type': 'string', 'description': '演示文稿总标题'},
                'slides': {'type': 'array', 'items': {'type': 'object', 'properties': {
                    'title': {'type': 'string'},
                    'bullets': {'type': 'array', 'items': {'type': 'string'}}
                }}},
            },
            'required': ['filename', 'title', 'slides'],
        },
        'handler': _h(lambda filename, title, slides: create_ppt(filename, title, slides)),
    },
    {
        'name': 'list_desktop_files',
        'description': '列出用户桌面上的文件（名称/大小/修改时间，最多 50 个）。',
        'inputSchema': {'type': 'object', 'properties': {'ext': {'type': 'string', 'description': '可选，按扩展名过滤，如 pdf'}}},
        'handler': _h(lambda ext=None: list_desktop_files(ext)),
    },
    {
        'name': 'open_file',
        'description': '用系统默认程序打开桌面或导出目录里的文件。',
        'inputSchema': {'type': 'object', 'properties': {
            'filename': {'type': 'string', 'description': '文件名（仅限桌面或导出目录内）'}},
            'required': ['filename']},
        'handler': _h(lambda filename: open_file(filename)),
    },
    {
        'name': 'fetch_webpage',
        'description': '抓取网页正文文本（去 HTML 标签，截断 8000 字符），用于联网查询资料。',
        'inputSchema': {'type': 'object', 'properties': {
            'url': {'type': 'string'}}, 'required': ['url']},
        'handler': _h(lambda url, max_chars=8000: fetch_webpage(url, max_chars)),
    },
]

def schemas():
    return [{'type': 'function', 'function': {'name': t['name'], 'description': t['description'],
                                              'parameters': t['inputSchema']}} for t in REGISTRY]

def describe():
    return [{'name': t['name'], 'description': t['description']} for t in REGISTRY]

def execute(name, args):
    for t in REGISTRY:
        if t['name'] == name:
            try:
                result = t['handler'](args or {})
                if isinstance(result, str):
                    return {'ok': True, 'text': result}
                return {'ok': True, 'data': result, 'text': json.dumps(result, ensure_ascii=False)[:2000]}
            except Exception as e:
                return {'ok': False, 'text': '工具执行失败：%s' % e}
    return {'ok': False, 'text': '未知工具：%s' % name}


# ================= v3.1 扩展装载（内建扩充 + 远程 MCP） =================
REMOTE_STATUS = []  # [{name, status, server, tool_count}] 供 /api/tools 展示


def load_extensions(mcp_conf_list=None, log=print):
    """启动时挂载：① 11 个新内建工具（tools_extra）② config.json mcpServers 远程 MCP。
    返回挂载的远程工具数。幂等保护：重复调用不会重复注册。"""
    import tools_extra
    existing = {t['name'] for t in REGISTRY}
    added = 0
    for t in tools_extra.builtin_tools():
        if t['name'] not in existing:
            REGISTRY.append(t)
            added += 1
    remote_count = 0
    for conf in (mcp_conf_list or []):
        name = conf.get('name') or 'mcp'
        if not (conf.get('url') or conf.get('command')):
            continue
        try:
            import mcp_client
            remote = mcp_client.load_remote_tools([conf], log=log)
        except Exception as e:
            log('MCP[%s] 装载失败：%s' % (name, str(e)[:200]))
            remote = []
        cli_tools = [t for t in remote if t.get('remote')]
        server_name = ''
        for t in remote:
            if t['name'] not in {x['name'] for x in REGISTRY}:
                REGISTRY.append({'name': t['name'], 'description': t['description'],
                                 'inputSchema': t['inputSchema'], 'handler': t['handler'], 'remote': True})
                remote_count += 1
        REMOTE_STATUS.append({'name': name, 'status': '已连接' if remote else '未连接',
                              'server': '', 'tool_count': len(cli_tools)})
    log('扩展装载完成：新增内建 %d 个，远程 MCP 工具 %d 个，注册表共 %d 个' % (added, remote_count, len(REGISTRY)))
    return remote_count
